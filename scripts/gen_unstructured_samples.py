# SPDX-License-Identifier: MIT
"""Regenerate the binary samples in samples/unstructured/ from the Markdown sources.

    uv run --group dev python scripts/gen_unstructured_samples.py

- 계약관리규정.pdf  <- 계약관리규정.md  (pymupdf Story; Korean system font AppleGothic)
- 청구수납규정.docx <- 청구수납규정.md  (python-docx; headings as Heading 1-3, GFM tables)

pymupdf and python-docx are dev-only dependencies: AutoOntology itself reads these files
with kordoc. Output is deterministic apart from the PDF/DOCX creation timestamps.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples" / "unstructured"
FONT = Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf")


def make_pdf(src: Path, out: Path) -> None:
    import pymupdf
    from markdown_it import MarkdownIt

    html = MarkdownIt("commonmark").enable("table").render(src.read_text(encoding="utf-8"))
    css = (
        "@font-face {font-family: kr; src: url(%s);} "
        "* {font-family: kr; font-size: 10.5pt; line-height: 1.5;} "
        "h1 {font-size: 18pt;} h2 {font-size: 14pt;} h3 {font-size: 12pt;} "
        "table {border-collapse: collapse;} td, th {border: 1px solid #444; padding: 3px;}"
        % FONT.name
    )
    story = pymupdf.Story(html=html, user_css=css, archive=pymupdf.Archive(str(FONT.parent)))
    writer = pymupdf.DocumentWriter(str(out))
    page = pymupdf.paper_rect("a4")
    where = page + (56, 56, -56, -56)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    doc = pymupdf.open(out)
    doc.set_metadata({"title": src.stem, "author": "AutoOntology samples"})
    doc.subset_fonts()  # the full AppleGothic face is ~15 MB
    tmp = out.with_suffix(".tmp.pdf")
    doc.save(tmp, garbage=4, deflate=True)
    doc.close()
    tmp.replace(out)


_TABLE_SEP = re.compile(r"^\|?\s*:?-{3,}")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def make_docx(src: Path, out: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = src.stem
    document.core_properties.author = "AutoOntology samples"
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif line.startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            header = _cells(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            table = document.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for r, values in enumerate([header, *rows]):
                for c, value in enumerate(values):
                    table.cell(r, c).text = value
            continue
        elif line.strip():
            document.add_paragraph(line)
        i += 1
    document.save(out)


def main() -> int:
    if not FONT.is_file():
        print(f"missing font {FONT}", file=sys.stderr)
        return 1
    make_pdf(SAMPLES / "계약관리규정.md", SAMPLES / "계약관리규정.pdf")
    make_docx(SAMPLES / "청구수납규정.md", SAMPLES / "청구수납규정.docx")
    for name in ("계약관리규정.pdf", "청구수납규정.docx"):
        print(f"wrote {SAMPLES / name} ({(SAMPLES / name).stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
