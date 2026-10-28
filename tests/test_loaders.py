# SPDX-License-Identifier: MIT
import json
import sqlite3

import duckdb
import pandas as pd
import pytest

from autoontology.structured import loaders
from autoontology.structured.loaders import (
    LoadError,
    dedupe_identifiers,
    discover_files,
    duckdb_to_pg_type,
    normalize_identifier,
    read_tables,
)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("﻿고객ID", "고객id"),
        ("Customer Name", "customer_name"),
        ("  계약 금액(원) ", "계약_금액_원"),
        ("a--b..c", "a_b_c"),
        ("___", "col"),
        ("", "col"),
        ("2024 매출", "2024_매출"),
    ],
)
def test_normalize_identifier(raw, expected):
    assert normalize_identifier(raw) == expected


def test_normalize_identifier_truncates_by_bytes():
    name = normalize_identifier("가" * 40)  # 120 bytes
    assert len(name.encode()) <= 63 and name == "가" * 21


def test_dedupe_identifiers():
    assert dedupe_identifiers(["a", "a", "b", "a"]) == ["a", "a_2", "b", "a_3"]
    long = "가" * 21
    out = dedupe_identifiers([long, long])
    assert out[1].endswith("_2") and len(out[1].encode()) <= 63


@pytest.mark.parametrize(
    ("duck", "pg"),
    [
        ("BIGINT", "bigint"),
        ("VARCHAR", "text"),
        ("DATE", "date"),
        ("DOUBLE", "double precision"),
        ("DECIMAL(10,2)", "numeric(10,2)"),
        ("BOOLEAN", "boolean"),
        ("TIMESTAMP WITH TIME ZONE", "timestamptz"),
        ("STRUCT(a INTEGER)", "jsonb"),
        ("INTEGER[]", "jsonb"),
        ("SOMETHING", "text"),
    ],
)
def test_duckdb_to_pg_type(duck, pg):
    con = duckdb.connect()
    if duck != "SOMETHING":
        typ = con.sql(f"SELECT NULL::{duck} AS x").types[0]
        assert duckdb_to_pg_type(typ) == pg
    assert duckdb_to_pg_type(duck) == pg


def _write_samples(d):
    (d / "고객 목록.csv").write_bytes(
        "﻿고객ID,고객 명,금액,등록일,금액\nC001,가나,1000,2024-01-02,3\nC002,다라,2500,2024-02-03,4\n".encode()
    )
    (d / "t.tsv").write_text("a\tb\n1\tx\n", encoding="utf-8")
    (d / "e.jsonl").write_text(
        json.dumps({"id": 1, "meta": {"k": "v"}})
        + "\n"
        + json.dumps({"id": 2, "meta": None})
        + "\n"
    )
    pd.DataFrame({"x": [1.5, 2.5]}).to_parquet(d / "p.parquet")
    with pd.ExcelWriter(d / "book.xlsx") as xw:
        pd.DataFrame({"이름": ["a", "b"], "값": [1, 2]}).to_excel(
            xw, sheet_name="시트1", index=False
        )
        pd.DataFrame({"mixed": [1, "two"]}).to_excel(xw, sheet_name="Other", index=False)
    with sqlite3.connect(d / "s.sqlite") as c:
        c.execute("create table Orders (id integer, Name text)")
        c.execute("insert into Orders values (1, 'x')")
    (d / "notes.md").write_text("ignored")
    sub = d / "sub"
    sub.mkdir()
    (sub / "deep.csv").write_text("a\n1\n")


def test_discover_files(tmp_path):
    _write_samples(tmp_path)
    names = [p.name for p in discover_files(tmp_path)]
    assert "notes.md" not in names and "deep.csv" not in names and len(names) == 6
    assert "deep.csv" in [p.name for p in discover_files(tmp_path, recursive=True)]
    with pytest.raises(LoadError):
        discover_files(tmp_path / "notes.md")
    with pytest.raises(LoadError):
        discover_files(tmp_path / "missing")


def test_read_tables_plans_types_and_names(tmp_path):
    _write_samples(tmp_path)
    duck = duckdb.connect()
    warnings: list[str] = []
    plans = {p.name: p for p in read_tables(discover_files(tmp_path), duck, warnings)}
    assert set(plans) == {"고객_목록", "t", "e", "p", "book_시트1", "book_other", "orders"}
    cust = plans["고객_목록"]
    assert [(c.name, c.pg_type) for c in cust.columns] == [
        ("고객id", "text"),
        ("고객_명", "text"),
        ("금액", "bigint"),
        ("등록일", "date"),
        ("금액_1", "bigint"),
    ]
    assert duck.execute(cust.select_sql).fetchall()[0][0] == "C001"
    assert dict((c.name, c.pg_type) for c in plans["e"].columns)["meta"] == "jsonb"
    meta = duck.execute(plans["e"].select_sql).fetchall()
    assert json.loads(meta[0][1]) == {"k": "v"} and meta[1][1] is None
    assert [c.pg_type for c in plans["book_other"].columns] == ["text"]
    assert [c.name for c in plans["orders"].columns] == ["id", "name"]
    assert plans["t"].columns[1].pg_type == "text"


def test_read_tables_dedupes_table_names(tmp_path):
    (tmp_path / "A.csv").write_text("x\n1\n")
    (tmp_path / "a.tsv").write_text("x\n1\n")
    warnings: list[str] = []
    plans = read_tables(sorted(tmp_path.iterdir()), duckdb.connect(), warnings)
    assert [p.name for p in plans] == ["a", "a_2"] and warnings


def test_load_requires_files(isolated_env):
    from autoontology.config import Settings

    with pytest.raises(LoadError):
        loaders.load_files_to_postgres([], "autoontology_src_x", Settings.from_env())
    with pytest.raises(ValueError):
        loaders.load_files_to_postgres([isolated_env], "other_db", Settings.from_env())


def test_sqlite_table_names_with_quotes_are_escaped(tmp_path):
    import sqlite3

    import duckdb

    from autoontology.structured.loaders import read_tables

    db = tmp_path / "q.sqlite"
    with sqlite3.connect(db) as conn:
        conn.execute('CREATE TABLE "we""ird" (x INTEGER)')
        conn.execute('INSERT INTO "we""ird" VALUES (7)')
        conn.execute('CREATE TABLE "weird" (y INTEGER)')  # what stripping would have hit
    plans = read_tables([db], duckdb.connect(), [])
    cols = {p.sheet: [c.name for c in p.columns] for p in plans}
    assert cols == {'we"ird': ["x"], "weird": ["y"]}


def test_duckdb_renames_duplicate_csv_headers_before_dedupe(tmp_path):
    import duckdb

    from autoontology.structured.loaders import read_tables

    path = tmp_path / "d.csv"
    path.write_text("금액,금액,Amount,amount,a b,a_b\n1,2,3,4,5,6\n", encoding="utf-8")
    (plan,) = read_tables([path], duckdb.connect(), [])
    assert [c.name for c in plan.columns] == [
        "금액",
        "금액_1",
        "amount",
        "amount_1",
        "a_b",
        "a_b_2",
    ]


def test_sqlite_nullable_integers_keep_precision(tmp_path):
    """Finding #3: an INTEGER column with NULLs must stay BIGINT, not float64."""
    db = tmp_path / "big.sqlite"
    with sqlite3.connect(db) as conn:
        conn.execute("CREATE TABLE t (id INTEGER, ref INTEGER)")
        conn.executemany(
            "INSERT INTO t VALUES (?, ?)", [(1, 9007199254740993), (2, None), (3, 9007199254740992)]
        )
    duck = duckdb.connect()
    (plan,) = read_tables([db], duck, [])
    assert {c.name: c.pg_type for c in plan.columns} == {"id": "bigint", "ref": "bigint"}
    rows = duck.execute(plan.select_sql).fetchall()
    assert rows == [(1, 9007199254740993), (2, None), (3, 9007199254740992)]


def test_manifest_and_diff(tmp_path):
    """Review #4/#5: the load manifest records each table's input and column types; the diff
    reports added, removed and changed (new column / changed type) tables."""
    from autoontology.structured.loaders import diff_manifest, table_manifest

    (tmp_path / "customers.csv").write_text("id,name\n1,a\n")
    (tmp_path / "orders.csv").write_text("id,total\n1,2.5\n")
    old = table_manifest(read_tables(discover_files(tmp_path), duckdb.connect(), []))
    assert old["customers"]["columns"] == {"id": "bigint", "name": "text"}
    assert old["orders"]["source_file"].endswith("orders.csv")
    (tmp_path / "orders.csv").unlink()
    (tmp_path / "customers.csv").write_text("id,name,loyalty_score\n1,a,0.5\n")
    (tmp_path / "items.csv").write_text("sku\nx\n")
    new = table_manifest(read_tables(discover_files(tmp_path), duckdb.connect(), []))
    assert diff_manifest(old, new) == {
        "added": ["items"],
        "removed": ["orders"],
        "changed": ["customers"],
    }
    assert diff_manifest(new, new) == {"added": [], "removed": [], "changed": []}
    retyped = {**new, "items": {**new["items"], "columns": {"sku": "bigint"}}}
    assert diff_manifest(new, retyped)["changed"] == ["items"]


def test_sqlite_blobs_stay_binary(tmp_path):
    """Fix wave 2 #1: nullable dtype inference turned BLOBs into text or failed on bytes
    that are not UTF-8; they must stay bytea, next to exact nullable BIGINTs."""
    db = tmp_path / "blob.sqlite"
    with sqlite3.connect(db) as conn:
        conn.execute("CREATE TABLE t (id INTEGER, payload BLOB, note TEXT)")
        conn.executemany(
            "INSERT INTO t VALUES (?, ?, ?)",
            [(9007199254740993, b"abc", "a"), (None, b"\xff\x00", None), (2, None, "c")],
        )
    duck = duckdb.connect()
    (plan,) = read_tables([db], duck, [])
    assert {c.name: c.pg_type for c in plan.columns} == {
        "id": "bigint",
        "payload": "bytea",
        "note": "text",
    }
    rows = duck.execute(plan.select_sql).fetchall()
    assert rows == [(9007199254740993, b"abc", "a"), (None, b"\xff\x00", None), (2, None, "c")]


def test_xlsx_keeps_text_cells_and_exact_integers(tmp_path):
    """Fix wave 2 #2: identifiers stored as text cells must not become doubles; integral
    numbers with blanks stay exact BIGINTs."""
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["고객ID", "count", "ratio", "active", "since", "mixed"])
    ws.append(["9007199254740993", 123456789012345, 1.5, True, dt.datetime(2024, 1, 2), 1])
    ws.append(["9007199254740992", None, 2, False, dt.datetime(2024, 1, 3), "x"])
    ws.append([None, 3, None, None, None, None])
    path = tmp_path / "ids.xlsx"
    wb.save(path)
    duck = duckdb.connect()
    (plan,) = read_tables([path], duck, [])
    assert {c.name: c.pg_type for c in plan.columns} == {
        "고객id": "text",
        "count": "bigint",
        "ratio": "double precision",
        "active": "boolean",
        "since": "timestamp",
        "mixed": "text",
    }
    rows = duck.execute(plan.select_sql).fetchall()
    assert [r[0] for r in rows] == ["9007199254740993", "9007199254740992", None]
    assert [r[1] for r in rows] == [123456789012345, None, 3]  # was float64 before
    assert [r[5] for r in rows] == ["1", "x", None]


def test_xlsx_out_of_range_numbers_and_dates_load_exactly(tmp_path):
    """Fix wave 3 #3: integers beyond int64 and dates outside pandas' nanosecond range
    (9999-12-31 sentinels, 1600-01-01) aborted the load; they must load exactly."""
    import datetime as dt

    import openpyxl

    from autoontology.structured.loaders import typed_column

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["big", "valid_to", "valid_from", "plain_date"])
    ws.append(
        [
            100000000000000000000,
            dt.datetime(9999, 12, 31),
            dt.datetime(1600, 1, 1),
            dt.datetime(2024, 1, 2),
        ]
    )
    ws.append([None, None, dt.datetime(2020, 5, 6), None])
    path = tmp_path / "ranges.xlsx"
    wb.save(path)
    duck = duckdb.connect()
    (plan,) = read_tables([path], duck, [])
    types = {c.name: c.pg_type for c in plan.columns}
    assert types["big"] == "numeric" and types["valid_to"] == "timestamp"
    assert types["valid_from"] == "timestamp" and types["plain_date"] == "timestamp"
    rows = duck.execute(plan.select_sql).fetchall()
    assert rows[0][0] == 100000000000000000000 and rows[1][0] is None
    assert rows[0][1] == dt.datetime(9999, 12, 31) and rows[1][1] is None
    assert [r[2] for r in rows] == [dt.datetime(1600, 1, 1), dt.datetime(2020, 5, 6)]
    # boundaries
    series, cast = typed_column([2**63 - 1, None])
    assert str(series.dtype) == "Int64" and cast is None
    series, cast = typed_column([2**63, None])
    assert cast == "HUGEINT" and series.tolist() == [str(2**63), None]
    series, cast = typed_column([2**200])
    assert cast == "VARCHAR" and series.tolist() == [str(2**200)]  # beyond HUGEINT: text
    series, cast = typed_column([dt.date(9999, 12, 31), None])
    assert cast == "DATE" and series.tolist() == ["9999-12-31", None]


def test_oversized_value_past_duckdbs_sample_keeps_the_planned_type(tmp_path):
    """Fix wave 4 #3: DuckDB types an object column from a sample (1000 rows); a value
    beyond int64 at row 1999 of 2001 must still plan numeric and load exactly."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "amount"])
    for i in range(2001):
        ws.append([i, 100000000000000000000 if i == 1999 else i])
    xlsx = tmp_path / "late.xlsx"
    wb.save(xlsx)
    db = tmp_path / "late.sqlite"
    blobs = [(None,)] * 1500 + [(b"\xff\x00",)] * 501  # BLOBs only after the sample
    with sqlite3.connect(db) as conn:
        conn.execute("CREATE TABLE b (payload BLOB)")
        conn.executemany("INSERT INTO b VALUES (?)", blobs)
    duck = duckdb.connect()
    plans = {p.name: p for p in read_tables([xlsx, db], duck, [])}
    assert {c.name: c.pg_type for c in plans["late"].columns} == {
        "id": "bigint",
        "amount": "numeric",
    }
    rows = duck.execute(plans["late"].select_sql).fetchall()
    assert rows[1999] == (1999, 100000000000000000000) and rows[5] == (5, 5)
    assert [c.pg_type for c in plans["b"].columns] == ["bytea"]
    assert duck.execute(plans["b"].select_sql).fetchall()[-1] == (b"\xff\x00",)


class _CopyConn:
    """Enough of a psycopg connection for write_tables: records DDL and COPY rows."""

    def __init__(self):
        self.sql, self.rows = [], {}

    def transaction(self):
        import contextlib

        return contextlib.nullcontext()

    def execute(self, sql, *a):
        self.sql.append(sql)

    def cursor(self):
        conn = self

        class _Cur:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def copy(self, sql):
                table = sql.split('"')[1]
                rows = conn.rows.setdefault(table, [])

                class _Copy:
                    def __enter__(self):
                        return self

                    def __exit__(self, *a):
                        return False

                    def write_row(self, row):
                        rows.append(tuple(row))

                return _Copy()

        return _Cur()


@pytest.mark.parametrize(
    ("headers", "names"),
    [
        ([1, "1", "keep"], ["1", "1_2", "keep"]),
        (["ID", "id", "keep"], ["id", "id_2", "keep"]),
    ],
)
def test_xlsx_colliding_headers_keep_every_column(tmp_path, headers, names):
    """Fix wave 5 #2: numeric header 1 vs text "1" lost a column, ID vs id (DuckDB is
    case-insensitive) copied one column's values into both."""
    import openpyxl

    from autoontology.structured.loaders import write_tables

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(["left", "right", "keep"])
    path = tmp_path / "hdr.xlsx"
    wb.save(path)
    duck = duckdb.connect()
    (plan,) = read_tables([path], duck, [])
    assert [c.name for c in plan.columns] == names
    assert [c.original for c in plan.columns] == [str(h) for h in headers]
    conn = _CopyConn()
    (report,) = write_tables([plan], duck, conn)
    assert report.rows == 1 and conn.rows["hdr"] == [("left", "right", "keep")]


def _write_xls(path, sheets):
    import datetime as dt

    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for name, rows in sheets.items():
        ws = wb.add_sheet(name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is None:
                    continue
                style = date_style if isinstance(value, dt.date) else xlwt.Style.default_style
                ws.write(r, c, value, style)
    wb.save(str(path))


def test_xls_one_table_per_sheet_with_types(tmp_path):
    import datetime as dt

    path = tmp_path / "suppliers.xls"
    _write_xls(
        path,
        {
            "거래처": [
                ("코드", "이름", "신용등급", "거래한도", "등록일"),
                ("S001", "한빛물산", 1, 1500.5, dt.date(2024, 1, 5)),
                ("S002", "동해상사", 2, 300.0, dt.date(2024, 3, 1)),
                ("S003", "서해무역", None, None, None),
            ],
            "등급": [("등급", "설명"), (1, "우수"), (2, "보통")],
            "빈시트": [],
        },
    )
    assert discover_files(path) == [path]
    assert path in discover_files(tmp_path)
    duck = duckdb.connect()
    warnings: list[str] = []
    plans = {p.name: p for p in read_tables([path], duck, warnings)}
    assert set(plans) == {"suppliers_거래처", "suppliers_등급"}
    assert any("빈시트" in w for w in warnings)
    sup = plans["suppliers_거래처"]
    assert sup.sheet == "거래처" and sup.source_file.endswith("suppliers.xls")
    assert {c.name: c.pg_type for c in sup.columns} == {
        "코드": "text",
        "이름": "text",
        "신용등급": "bigint",
        "거래한도": "double precision",
        "등록일": "date",
    }
    rows = duck.execute(sup.select_sql).fetchall()
    assert rows[0] == ("S001", "한빛물산", 1, 1500.5, dt.date(2024, 1, 5))
    assert rows[2] == ("S003", "서해무역", None, None, None)
    assert [c.pg_type for c in plans["suppliers_등급"].columns] == ["bigint", "text"]


def test_xls_single_sheet_is_named_after_the_file(tmp_path):
    path = tmp_path / "단가표.xls"
    _write_xls(path, {"Sheet1": [("품목", "단가"), ("볼트", 120)]})
    (plan,) = read_tables([path], duckdb.connect(), [])
    assert plan.name == "단가표" and plan.sheet == "Sheet1"
