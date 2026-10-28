"""Loaders against a real, temporary Postgres database (integration marker)."""

import datetime as dt

import pytest

from autoontology.config import PG_MAINTENANCE_DB, Settings
from autoontology.engine import pg
from autoontology.structured.loaders import discover_files, load_files_to_postgres

pytestmark = pytest.mark.integration

DB = "autoontology_src_pytest_loader"


@pytest.fixture
def settings():
    s = Settings.from_env()
    yield s
    with pg.connect(s, PG_MAINTENANCE_DB) as conn:
        conn.execute(f'DROP DATABASE IF EXISTS "{DB}" WITH (FORCE)')


def test_load_and_replace(tmp_path, settings):
    f = tmp_path / "계약 목록.csv"
    f.write_bytes("﻿계약ID,금액,시작일,메모\nK1,100,2024-01-01,가\nK2,200,2024-02-01,\n".encode())
    (tmp_path / "e.jsonl").write_text('{"id": 1, "tags": ["a", "b"]}\n')
    report = load_files_to_postgres(discover_files(tmp_path), DB, settings)
    assert report.created_database
    by_name = {t.name: t for t in report.tables}
    assert by_name["계약_목록"].rows == 2
    with pg.connect(settings, DB) as conn:
        rows = conn.execute('select "계약id", "금액", "시작일", "메모" from "계약_목록" order by 1').fetchall()
        assert rows == [("K1", 100, dt.date(2024, 1, 1), "가"), ("K2", 200, dt.date(2024, 2, 1), None)]
        types = dict(conn.execute(
            "select column_name, data_type from information_schema.columns "
            "where table_name = '계약_목록'").fetchall())
        assert types == {"계약id": "text", "금액": "bigint", "시작일": "date", "메모": "text"}
        assert conn.execute('select tags from "e"').fetchone()[0] == ["a", "b"]
    # Replace: a second load with fewer rows replaces the table; e.jsonl is no longer an
    # input, so its table is dropped (review #5); the column change is reported before any
    # write (review #4).
    f.write_text("계약ID,금액\nK9,1\n", encoding="utf-8")
    seen = []

    def on_changes(changes):
        with pg.connect(settings, DB) as conn:  # still the old schema at this point
            cols = conn.execute("select count(*) from information_schema.columns "
                                "where table_name='계약_목록'").fetchone()[0]
        seen.append((changes, cols))

    report = load_files_to_postgres([f], DB, settings, on_changes=on_changes)
    assert not report.created_database and report.tables[0].rows == 1
    assert seen == [({"added": [], "removed": ["e"], "changed": ["계약_목록"]}, 4)]
    assert report.dropped_tables == ["e"]
    with pg.connect(settings, DB) as conn:
        assert conn.execute('select count(*) from "계약_목록"').fetchone()[0] == 1
        assert conn.execute(
            "select count(*) from information_schema.columns where table_name='계약_목록'"
        ).fetchone()[0] == 2


def test_xlsx_and_sqlite_edge_values_reach_postgres_exactly(tmp_path, settings):
    """Fix waves 2/3: text ids, beyond-int64 integers, 9999/1600 dates, BLOBs and NULLs
    survive DuckDB -> COPY -> Postgres."""
    import sqlite3

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id_text", "big", "valid_to", "valid_from"])
    ws.append(["9007199254740993", 100000000000000000000, dt.datetime(9999, 12, 31),
               dt.datetime(1600, 1, 1)])
    ws.append([None, None, None, dt.datetime(2020, 5, 6)])
    wb.save(tmp_path / "edge.xlsx")
    with sqlite3.connect(tmp_path / "bin.sqlite") as conn:
        conn.execute("CREATE TABLE blobs (id INTEGER, payload BLOB)")
        conn.executemany("INSERT INTO blobs VALUES (?, ?)",
                         [(9007199254740993, b"\xff\x00"), (None, None), (2, b"abc")])
    load_files_to_postgres(discover_files(tmp_path), DB, settings)
    with pg.connect(settings, DB) as conn:
        rows = conn.execute('select id_text, big, valid_to, valid_from from "edge" '
                            "order by valid_from").fetchall()
        assert rows == [("9007199254740993", 100000000000000000000, dt.datetime(9999, 12, 31),
                         dt.datetime(1600, 1, 1)),
                        (None, None, None, dt.datetime(2020, 5, 6))]
        blobs = conn.execute('select id, payload from "blobs" order by id nulls last').fetchall()
        assert [(r[0], bytes(r[1]) if r[1] is not None else None) for r in blobs] == [
            (2, b"abc"), (9007199254740993, b"\xff\x00"), (None, None)]
        types = dict(conn.execute("select column_name, data_type from information_schema.columns "
                                  "where table_name = 'edge'").fetchall())
        assert types == {"id_text": "text", "big": "numeric",
                         "valid_to": "timestamp without time zone",
                         "valid_from": "timestamp without time zone"}


def test_late_oversized_integer_reaches_postgres_as_numeric(tmp_path, settings):
    """Fix wave 4 #3: 1e20 at row 1999 of 2001 (past DuckDB's 1000-row sample)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "amount"])
    for i in range(2001):
        ws.append([i, 100000000000000000000 if i == 1999 else i])
    wb.save(tmp_path / "late.xlsx")
    load_files_to_postgres(discover_files(tmp_path), DB, settings)
    with pg.connect(settings, DB) as conn:
        kind = conn.execute("select data_type from information_schema.columns where "
                            "table_name = 'late' and column_name = 'amount'").fetchone()[0]
        assert kind == "numeric"
        assert conn.execute('select amount from "late" where id = 1999').fetchone()[0] == \
            100000000000000000000
        assert conn.execute('select count(*) from "late"').fetchone()[0] == 2001
